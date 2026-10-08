"""Synthetic amount changes exercise construction, not memorized benchmark codes."""
import copy
import importlib.util
import json
import io
from decimal import Decimal
from pathlib import Path
import unittest
from unittest.mock import patch

SPEC = importlib.util.spec_from_file_location('estimate_calculation', Path(__file__).parents[1] / 'tools/estimate_calculation.py')
M = importlib.util.module_from_spec(SPEC); SPEC.loader.exec_module(M)


def request(amount='200'):
    return {'mode': 'BASE_ONLY', 'scope': 'Synthetic source-confirmed work only; no real project approval', 'price_basis': 'BASE_NET', 'vat': {'percent': '0', 'already_included': False, 'evidence': 'synthetic tax exemption', 'applicability_confirmed': True}, 'lines': [
        {'trace_id': 'VOR:revision-X:17', 'description': 'Synthetic unknown work',
         'quantity': {'value': amount, 'mode': 'physical', 'physical_unit': 'm2', 'norm_factor': '100', 'evidence': 'confirmed synthetic quantity'},
         'norm': {'code': 'TEST-UNSEEN-42', 'revision': 'test-only', 'unit': '100m2', 'physical_unit': 'm2', 'norm_factor': '100', 'source': 'synthetic fixture, never production', 'evidence': 'fixture', 'applicability_confirmed': True,
                  'components': {'labour': '10', 'machines': '20', 'machinist': '5', 'materials': '30'},
                  'resources': [{'id': 'included-X', 'unit': 'kg', 'quantity_per_norm': '3', 'included': True, 'evidence': 'fixture'}, {'id': 'excluded-Y', 'unit': 'piece', 'quantity_per_norm': '2', 'included': False, 'evidence': 'fixture', 'external_price': {'value': '7', 'unit': 'piece', 'basis': 'BASE_NET', 'evidence': 'fixture price'}}]},
         'coefficients': [], 'indexes': [], 'no_extra_coefficients': {'evidence': 'Synthetic server-reviewed no-extra-coefficients determination'},
         'overhead': {'percent': '100', 'basis': 'FOT', 'scope_trace_id': 'VOR:revision-X:17', 'evidence': 'fixture', 'applicability_confirmed': True},
         'profit': {'percent': '50', 'basis': 'FOT', 'scope_trace_id': 'VOR:revision-X:17', 'evidence': 'fixture', 'applicability_confirmed': True}}]}


class CalculationTests(unittest.TestCase):
    def test_changed_unseen_amount_and_code(self):
        a = M.build_estimate(request('200')); b = M.build_estimate(request('300'))
        self.assertEqual(a['totals']['net'], '193.00')
        self.assertEqual(b['totals']['net'], '289.50')
        self.assertEqual(a['lines'][0]['resources'][0]['quantity'], '6')
        self.assertEqual(a['approval'], 'NOT_FOR_APPROVAL')

    def test_explicit_normalized_mode_no_second_division(self):
        q = {'value': '2', 'mode': 'norm_units', 'physical_unit': 'm2', 'norm_factor': '100', 'trace_id': 'Q', 'evidence': 'test'}
        r = M.convert_quantity(q)
        self.assertEqual(r['physical'], '200'); self.assertEqual(r['norm_units'], '2')

    def test_semantic_dimensions_not_first_number(self):
        q = {'value': '40', 'mode': 'geometry', 'physical_unit': 'piece', 'norm_factor': '1', 'trace_id': 'Q', 'evidence': 'passport', 'relation_confirmed': True,
             'input_unit': 'm', 'relation': 'linear_to_pieces', 'geometry': {'length': {'semantic': 'length', 'value': '2000', 'unit': 'mm', 'evidence': 'passport length'}}}
        self.assertEqual(M.convert_quantity(q)['physical'], '2E+1')
        q['geometry']['length']['semantic'] = 'width'
        with self.assertRaises(ValueError): M.convert_quantity(q)
        q['geometry']['length']['semantic'] = 'length'; q['relation_confirmed'] = False
        with self.assertRaises(ValueError): M.convert_quantity(q)

    def test_included_resource_double_count_rejected(self):
        r = request(); r['lines'][0]['norm']['resources'][0]['external_price'] = {'value': '99'}
        with self.assertRaises(ValueError): M.build_estimate(r)

    def test_double_vat_rejected(self):
        r = request(); r['vat']['already_included'] = True
        with self.assertRaises(ValueError): M.build_estimate(r)
        r = request(); r['price_basis'] = 'CURRENT_GROSS'
        with self.assertRaises(ValueError): M.build_estimate(r)

    def test_missing_is_not_zero(self):
        r = request(); r['lines'][0]['norm']['components']['materials'] = None
        result = M.build_estimate(r)
        self.assertIsNone(result['totals']['net']); self.assertTrue(result['missing_data'])
        r = request(); r['lines'][0]['norm']['resources'][1]['external_price'] = None
        self.assertIsNone(M.build_estimate(r)['totals']['net'])
        with self.assertRaises(ValueError): M.decimal(None)
        self.assertEqual(M.decimal('0'), 0)

    def test_negative_cost_effect_and_missing(self):
        r = {'trace_id': 'scenario-Q', 'resource_id': 'same-X', 'unit': 'm3', 'old_unit': 'm3', 'new_unit': 'm3', 'same_resource': True, 'source': 'synthetic source',
             'old_quantity': '50', 'new_quantity': '20', 'per_unit_components': {'materials': '7.1'}, 'price_basis': 'CURRENT_NET', 'price_evidence': 'exact price', 'scope': 'Material-only scenario; no labour/NR/SP/VAT effect asserted', 'conditions': ['Same scoped resource and unit']}
        result = M.cost_effect(r)
        self.assertEqual(result['quantity_delta'], '-30'); self.assertEqual(result['cost_effect_net'], '-213.00')
        r['per_unit_components']['materials'] = None
        self.assertIsNone(M.cost_effect(r)['cost_effect_net'])
        r['per_unit_components']['materials'] = '7.1'; r['new_unit'] = 'piece'
        with self.assertRaises(ValueError): M.cost_effect(r)

    def test_project_defined_resource_unknown_and_confirmed_override(self):
        for sentinel in (None, 'П'):
            r = request(); resource = r['lines'][0]['norm']['resources'][1]
            resource['quantity_per_norm'] = sentinel
            result = M.build_estimate(r)
            self.assertIsNone(result['lines'][0]['resources'][1]['quantity'])
            self.assertIsNone(result['totals']['net'])
            self.assertTrue(any('confirmed project consumption' in issue for issue in result['missing_data']))
            resource['project_quantity'] = {'value': '4', 'unit': 'piece', 'confirmed': True, 'evidence': 'confirmed project consumption sheet'}
            result = M.build_estimate(r)
            self.assertEqual(result['lines'][0]['resources'][1]['quantity'], '4')
            self.assertEqual(result['totals']['net'], '193.00')
            resource['project_quantity']['confirmed'] = False
            with self.assertRaises(ValueError): M.build_estimate(r)

    def test_scoped_indexes_and_missing_applicability(self):
        r = request(); r['mode'] = 'PREVIEW'
        r['lines'][0]['indexes'] = [{'component': c, 'value': '2', 'scope_trace_id': 'VOR:revision-X:17', 'evidence': 'fixture index', 'applicability_confirmed': True} for c in M.COMPONENTS]
        self.assertEqual(M.build_estimate(r, trusted_resolution=lambda *args: True)['totals']['net'], '386.00')
        r['lines'][0]['norm']['applicability_confirmed'] = False
        result = M.build_estimate(r)
        self.assertEqual(result['normative_applicability'], 'UNCONFIRMED')
        self.assertEqual(result['approval'], 'NOT_FOR_APPROVAL')
        r['lines'][0]['indexes'][0]['scope_trace_id'] = 'wrong-line'
        with self.assertRaises(ValueError): M.build_estimate(r)

    def preview_request(self):
        r = request(); r['mode'] = 'PREVIEW'
        r['lines'][0]['indexes'] = [{'component': c, 'value': '2', 'scope_trace_id': 'VOR:revision-X:17', 'evidence': 'fixture index', 'applicability_confirmed': True} for c in M.COMPONENTS]
        return r

    def test_forged_json_booleans_cannot_authorize_current_costs(self):
        r = self.preview_request()
        r['trusted_resolution'] = True
        result = M.build_estimate(r)
        self.assertIsNone(result['totals']['net'])
        self.assertIsNone(result['totals']['gross'])
        self.assertTrue(result['lines'][0]['current_costs_blocked'])
        self.assertTrue(all(v is None for v in result['lines'][0]['components_exact'].values()))
        self.assertTrue(any('server authority' in question for question in result['missing_data']))

    def test_server_resolution_binds_trace_evidence_and_exact_facts(self):
        r = self.preview_request(); ledger = set()
        def key(kind, trace, evidence, facts):
            return json.dumps([kind, trace, evidence, facts], ensure_ascii=False, sort_keys=True)
        def synthetic_server_capture(kind, trace, evidence, facts):
            ledger.add(key(kind, trace, evidence, facts)); return True
        M.build_estimate(r, trusted_resolution=synthetic_server_capture)
        def synthetic_server_resolve(kind, trace, evidence, facts):
            return key(kind, trace, evidence, facts) in ledger
        valid = M.build_estimate(r, trusted_resolution=synthetic_server_resolve)
        self.assertEqual(valid['totals']['net'], '386.00')
        self.assertEqual(valid['normative_applicability'], 'SERVER_RESOLVED_FOR_DRAFT')
        for mutation in ('evidence', 'value', 'copied_trace', 'norm_source'):
            changed = copy.deepcopy(r)
            if mutation == 'evidence': changed['lines'][0]['indexes'][0]['evidence'] = 'different evidence'
            elif mutation == 'value': changed['lines'][0]['indexes'][0]['value'] = '99'
            elif mutation == 'norm_source': changed['lines'][0]['norm']['source'] = 'forged norm identity'
            else:
                line = changed['lines'][0]; line['trace_id'] = 'COPIED-LINE'
                for rule in line['indexes'] + [line['overhead'], line['profit']]: rule['scope_trace_id'] = 'COPIED-LINE'
            with self.subTest(mutation=mutation):
                result = M.build_estimate(changed, trusted_resolution=synthetic_server_resolve)
                self.assertIsNone(result['totals']['net'])
                self.assertTrue(result['lines'][0]['current_costs_blocked'])
        changed = copy.deepcopy(r); changed['lines'][0]['norm']['applicability_confirmed'] = False
        self.assertEqual(M.build_estimate(changed, trusted_resolution=synthetic_server_resolve)['totals']['net'], '386.00')

    def test_empty_coefficients_need_explicit_server_no_extra_determination(self):
        r = self.preview_request(); r['lines'][0].pop('no_extra_coefficients')
        result = M.build_estimate(r, trusted_resolution=lambda *args: True)
        self.assertIsNone(result['totals']['net'])
        self.assertTrue(any('no-extra-coefficients' in question for question in result['missing_data']))
        r = self.preview_request()
        result = M.build_estimate(r, trusted_resolution=lambda kind, *args: kind != 'profit')
        self.assertIsNone(result['totals']['net'])

    def test_copied_coefficient_evidence_cannot_apply_rule_twice(self):
        r = self.preview_request()
        coefficient = {'component': 'labour', 'value': '1.2', 'scope_trace_id': 'VOR:revision-X:17', 'evidence': 'one scoped source coefficient'}
        r['lines'][0]['coefficients'] = [coefficient, dict(coefficient, value='1.20', applicability_confirmed=True)]
        with self.assertRaisesRegex(ValueError, 'same rule twice'):
            M.build_estimate(r, trusted_resolution=lambda *args: True)

    def test_displayed_line_rounding_reconciles_net_vat_gross(self):
        r = request('100'); r['lines'][0]['norm']['resources'] = []
        r['lines'][0]['norm']['components'] = {'labour': '0.005', 'machines': '0', 'machinist': '0', 'materials': '0'}
        r['lines'][0]['overhead']['percent'] = '100'; r['lines'][0]['profit']['percent'] = '100'
        second = copy.deepcopy(r['lines'][0]); second['trace_id'] = 'VOR:revision-X:18'
        for field in ('overhead', 'profit'): second[field]['scope_trace_id'] = second['trace_id']
        r['lines'].append(second); r['vat']['percent'] = '20'
        result = M.build_estimate(r)
        self.assertEqual([line['direct'] for line in result['lines']], ['0.01', '0.01'])
        self.assertEqual([line['net'] for line in result['lines']], ['0.03', '0.03'])
        self.assertEqual(result['totals'], {'net': '0.06', 'vat': '0.01', 'gross': '0.07'})
        self.assertEqual(result['diagnostics']['unrounded_net_exact'], '0.030')
        self.assertTrue(all(line['net_exact'] == '0.015' for line in result['lines']))

    def exported_workbook(self, result):
        from openpyxl import Workbook, load_workbook
        # Real workbook serialization to memory, no managed temporary files.
        buffer = io.BytesIO(); save = Workbook.save
        with patch.object(Workbook, 'save', lambda book, path: save(book, buffer)):
            M.export_xlsx(result, '/ai-data/research/estimates/reference-07-01-01/IN_MEMORY_TEST_ONLY.xlsx')
        buffer.seek(0)
        return load_workbook(buffer, data_only=False)

    def test_xlsx_numeric_quantities_prices_totals_and_formula_injection(self):
        r = request('200'); r['lines'][0]['description'] = '=cmd'; r['lines'][0]['norm']['source'] = '=HYPERLINK("https://example.invalid","source")'
        second = copy.deepcopy(r['lines'][0]); second['trace_id'] = 'VOR:revision-X:18'; second['quantity']['value'] = '100'
        for field in ('overhead', 'profit'): second[field]['scope_trace_id'] = second['trace_id']
        r['lines'].append(second); result = M.build_estimate(r)
        book = self.exported_workbook(result); sheet = book.active
        for row in (5, 6):
            for column in (5, 7, 8, 9, 10, 11):
                self.assertEqual(sheet.cell(row, column).data_type, 'n')
                self.assertIsInstance(sheet.cell(row, column).value, (int, float))
            for column in (8, 9, 10, 11): self.assertEqual(sheet.cell(row, column).number_format, '0.00')
            self.assertEqual(sheet.cell(row, 2).value, '=cmd')
            self.assertEqual(sheet.cell(row, 2).data_type, 's')
            self.assertEqual(sheet.cell(row, 12).data_type, 's')
        net_sum = sum((Decimal(str(sheet.cell(row, 11).value)) for row in (5, 6)), Decimal(0))
        self.assertEqual(net_sum, Decimal(str(sheet['B7'].value)))
        self.assertEqual(net_sum, Decimal(result['totals']['net']))
        self.assertEqual(sheet['B7'].data_type, 'n'); self.assertEqual(sheet['B7'].number_format, '0.00')
        book.close()

    def test_xlsx_unknown_cost_remains_visible_text(self):
        result = M.build_estimate(self.preview_request())
        book = self.exported_workbook(result); sheet = book.active
        self.assertEqual(sheet['K5'].value, 'НЕИЗВЕСТНО')
        self.assertEqual(sheet['K5'].data_type, 's')
        self.assertEqual(sheet['B6'].value, 'НЕИЗВЕСТНО')
        self.assertEqual(sheet['E5'].data_type, 'n')
        book.close()


if __name__ == '__main__': unittest.main()
