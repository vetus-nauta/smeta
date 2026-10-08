import io
import sys
import unittest
from pathlib import Path
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'tools'))
from estimate_audit import inspect_file
from estimate_vor import extract_vor

def inspection(book):
 raw=io.BytesIO();book.save(raw);return {'files':[inspect_file('source.xlsx',raw.getvalue())]}

def parameters(book,scope='Base direct cost scenario only'):
 ws=book.create_sheet('Параметры')
 for row in [('namespace','FSNB-2020/ФЕР'),('revision','2021-12-20-d9'),('scope',scope),('base_direct_only',True)]:ws.append(row)

class VorTests(unittest.TestCase):
 def test_multisheet_three_sections_repeated_headers_all45(self):
  b=Workbook();s=b.active;s.title='ВОР часть1'
  s.append(['Раздел 1. Земляные работы']);s.merge_cells('A1:E1')
  s.append(['№ п/п',None,'Наименование и характеристика работ','Ед. изм.','Кол-во (объем работ)'])
  for n in range(15):s.append([n+1,None,'Разработка грунта','м3',n+10])
  s.append([None,None,'Итого по разделу',None,999])
  s.append([None,None,'Раздел 2. Основания']);s.merge_cells('C19:E19')
  s.append(['№ п/п',None,'Наименование работ','Единица измерения','Количество'])
  for n in range(15):s.append([n+1,None,'Устройство слоя','100 м3','1,5'])
  t=b.create_sheet('ВОР часть2');t.append(['№','Описание работ','Ед. изм.','Количество'])
  t.append([None,'Раздел 3. Покрытия']);t.merge_cells('B2:D2')
  for n in range(15):t.append([n+1,'Устройство покрытия','м2',n+1])
  parameters(b)
  r=extract_vor(inspection(b));self.assertEqual(len(r['lines']),45);self.assertEqual(len(r['source_sections']),3)
  self.assertEqual(len({x['id'] for x in r['lines']}),45);self.assertEqual(r['source_work_rows'],45)
  self.assertEqual(r['lines'][15]['quantity']['value'],'150.0');self.assertEqual(r['lines'][15]['quantity']['physical_unit'],'m3')
  self.assertEqual(r['lines'][15]['quantity']['evidence']['explicit_source_unit_factor'],'100')
  self.assertEqual(r['normative_scope']['revision'],'2021-12-20-d9');self.assertFalse(r['lines'][0]['overhead']['applicability_confirmed'])
  self.assertTrue(any(x['kind']=='TOTAL' for x in r['source_row_registry']))
 def test_unknown_quantity_unit_and_unverified_formula_reported(self):
  b=Workbook();s=b.active;s.append(['Наименование работ','Ед. изм.','Количество'])
  s.append(['Работа без числа','м3',None]);s.append(['Работа без единицы',None,12]);s.append(['Формула','м3','=2+3']);s.append(['Две единицы','м2/м3',2]);s.append(['Верная работа','м3',0])
  r=extract_vor(inspection(b));self.assertEqual(len(r['lines']),1);self.assertEqual(r['lines'][0]['quantity']['value'],'0');self.assertEqual(len(r['rejected_rows']),4)
  self.assertTrue(any('formula' in x['reason'] for x in r['rejected_rows']));self.assertTrue(any('empty' in x['reason'] for x in r['rejected_rows']))
 def test_no_header_no_quantity_inference(self):
  b=Workbook();b.active.append(['Устройство слоя','м3',77]);r=extract_vor(inspection(b));self.assertEqual(r['lines'],[]);self.assertTrue(r['rejected_rows']);self.assertTrue(any('header' in m for m in r['missing_data']))
 def test_ambiguous_duplicate_header_stays_rejected(self):
  b=Workbook();s=b.active;s.append(['Работа','Ед. изм.','Количество','Кол-во']);s.append(['Монтаж','шт',7,8]);r=extract_vor(inspection(b));self.assertEqual(r['lines'],[]);self.assertTrue(any(x['kind']=='HEADER' for x in r['rejected_rows']))
 def test_multirow_numbered_header_is_not_work_or_gap(self):
  b=Workbook();s=b.active;s.append(['№ п/п',None,'Наименование работ','Ед. изм.','Количество'])
  s.append([1,2,3,4,5]);s.append([1,None,'Устройство слоя','м3',7])
  r=extract_vor(inspection(b));self.assertEqual(len(r['lines']),1);self.assertEqual(r['source_work_rows'],1)
  self.assertFalse(any(x['row']==2 for x in r['rejected_rows']));self.assertTrue(any(x['kind']=='HEADER_NUMBERS' for x in r['source_row_registry']))
 def test_bool_flag_without_explicit_scope_is_not_authority(self):
  b=Workbook();s=b.active;s.append(['Работа','unit','quantity']);s.append(['Монтаж','шт',2]);parameters(b,'General source')
  r=extract_vor(inspection(b));self.assertNotIn('overhead',r['lines'][0]);self.assertFalse(r['normative_scope']['applicability_confirmed'])
 def test_matched_formula_only_and_physical_source_once(self):
  b=Workbook();s=b.active;s.append(['Работа','unit','quantity']);s.append(['Монтаж','100 м2','=2+3']);data=inspection(b);ws=data['files'][0]['sheets'][0]
  next(c for c in ws['cells'] if c['cell']=='C2')['cached']=5
  next(f for f in ws['formulas'] if f['cell']=='C2')['status']='MATCH'
  r=extract_vor(data);self.assertEqual(r['lines'][0]['quantity']['value'],'500');self.assertEqual(r['lines'][0]['quantity']['mode'],'physical')
 def test_metadata_conflict_and_no_latest_substitution(self):
  b=Workbook();s=b.active;s.append(['Работа','unit','quantity']);s.append(['Монтаж','шт',2]);parameters(b);t=b.create_sheet('Metadata');t.append(['revision','different'])
  r=extract_vor(inspection(b));self.assertIsNone(r['normative_scope']['revision']);self.assertTrue(any('Conflicting' in m for m in r['missing_data']))
 def test_exceeds5000_work_rows_fails_instead_of_truncating(self):
  b=Workbook();s=b.active;s.append(['Работа','unit','quantity'])
  for n in range(5001):s.append(['Работа','м3',1])
  with self.assertRaisesRegex(ValueError,'5000'):extract_vor(inspection(b))

if __name__=='__main__':unittest.main()
