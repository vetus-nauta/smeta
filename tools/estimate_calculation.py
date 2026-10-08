"""Deterministic construction/scenarios for professional review, never approval.

All prices are net BASE rubles per normative unit. Included resources are
already in norm.components.materials. External prices may only cover explicitly
excluded resources. No normative catalogue, rate or index is invented here.
"""
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import argparse
import json

COMPONENTS = ('labour', 'machines', 'machinist', 'materials')
DIMENSIONS = {'length', 'width', 'height'}


def decimal(value, *, nonnegative=True):
    if value is None or isinstance(value, bool):
        raise ValueError('Missing or boolean number; unknown is not zero')
    try:
        number = Decimal(str(value).replace(',', '.'))
    except InvalidOperation as exc:
        raise ValueError('Invalid decimal') from exc
    if not number.is_finite() or (nonnegative and number < 0):
        raise ValueError('Finite nonnegative number required')
    return number


def evidence(value):
    if not isinstance(value, (str, dict)) or not value:
        raise ValueError('Explicit source evidence required')
    return value


def convert_quantity(request):
    """Explicit normalization; optional geometry must name its semantic axes."""
    mode = request.get('mode')
    factor = decimal(request.get('norm_factor'))
    if not factor:
        raise ValueError('Positive norm_factor required')
    unit = request.get('physical_unit')
    if not isinstance(unit, str) or not unit:
        raise ValueError('Explicit physical_unit required')
    trace = request.get('trace_id')
    if not isinstance(trace, str) or not trace:
        raise ValueError('trace_id required')
    geometry = request.get('geometry')
    if geometry is not None:
        if mode != 'geometry':
            raise ValueError('Geometry requires explicit geometry mode')
        if request.get('relation_confirmed') is not True:
            raise ValueError('Geometry semantics must be confirmed')
        relation = request.get('relation')
        expected = {'linear_to_pieces': {'length'},
                    'area_to_volume': {'height'},
                    'pieces_to_volume': DIMENSIONS}.get(relation)
        if expected is None or set(geometry) != expected:
            raise ValueError('Explicit semantic length/width/height mapping required')
        sizes = {}
        for axis, dimension in geometry.items():
            if not isinstance(dimension, dict) or dimension.get('semantic') != axis:
                raise ValueError('Dimension label and semantic axis disagree')
            evidence(dimension.get('evidence'))
            scale = {'m': Decimal(1), 'cm': Decimal('.01'), 'mm': Decimal('.001')}.get(dimension.get('unit'))
            if scale is None:
                raise ValueError('Geometry unit must be m/cm/mm')
            sizes[axis] = decimal(dimension.get('value')) * scale
            if not sizes[axis]:
                raise ValueError('Geometry dimension must be positive')
        amount = decimal(request.get('value'))
        input_unit = request.get('input_unit')
        if relation == 'linear_to_pieces':
            if input_unit != 'm' or unit != 'piece':
                raise ValueError('linear_to_pieces requires m -> piece')
            physical = amount / sizes['length']
        elif relation == 'area_to_volume':
            if input_unit != 'm2' or unit != 'm3':
                raise ValueError('area_to_volume requires m2 -> m3')
            physical = amount * sizes['height']
        else:
            if input_unit != 'piece' or unit != 'm3':
                raise ValueError('pieces_to_volume requires piece -> m3')
            physical = amount * sizes['length'] * sizes['width'] * sizes['height']
    elif mode == 'physical':
        physical = decimal(request.get('value'))
    elif mode == 'norm_units':
        physical = decimal(request.get('value')) * factor
    else:
        raise ValueError('Explicit physical/norm_units/geometry mode required')
    evidence(request.get('evidence'))
    return {'trace_id': trace, 'physical': str(physical), 'norm_units': str(physical / factor),
            'physical_unit': unit, 'norm_factor': str(factor), 'mode': mode,
            'evidence': request['evidence'], 'relation_confirmed': request.get('relation_confirmed', None)}


def rounded(value):
    return str(value.quantize(Decimal('.01'), rounding=ROUND_HALF_UP)) if value is not None else None


def resolved_authority(resolver, kind, trace, source_evidence, facts):
    """Only server-injected code can resolve authority; JSON booleans cannot.

    facts are a detached JSON copy, so a callback cannot mutate calculation
    inputs. A resolver must bind kind/trace/evidence/facts to its trusted ledger.
    """
    if not callable(resolver):
        return False
    try:
        detached = json.loads(json.dumps(facts, ensure_ascii=False, sort_keys=True))
        detached_evidence = json.loads(json.dumps(source_evidence, ensure_ascii=False, sort_keys=True))
        return resolver(kind, trace, detached_evidence, detached) is True
    except Exception:
        return False


def authority_facts(value):
    return {k: v for k, v in value.items() if k != 'applicability_confirmed'}


def apply_factors(values, factors, trace, missing, *, resolver=None, kind='coefficient', scope=None):
    result = dict(values)
    seen_rules = set()
    for i, item in enumerate(factors):
        if item.get('scope_trace_id') != trace:
            raise ValueError('Coefficient/index scope must match trace_id')
        component = item.get('component')
        if component not in COMPONENTS:
            raise ValueError('Unknown component scope')
        evidence(item.get('evidence'))
        factor = decimal(item.get('value'))
        if factor <= 0:
            raise ValueError('Coefficient/index must be positive')
        identity = json.dumps([component, str(factor.normalize()), item['evidence']], ensure_ascii=False, sort_keys=True)
        if identity in seen_rules:
            raise ValueError('Duplicate scoped factor evidence would apply the same rule twice')
        seen_rules.add(identity)
        if not resolved_authority(resolver, kind, trace, item['evidence'], {'rule': authority_facts(item), 'scope': scope}):
            missing.append(f'{trace}: {kind}[{i}] applicability not resolved by server authority')
        if result[component] is not None:
            result[component] *= factor
    return result


def total(values):
    # Machinist remuneration is a subset of machines, never added twice.
    return None if any(values[c] is None for c in ('labour', 'machines', 'materials')) else sum((values[c] for c in ('labour', 'machines', 'materials')), Decimal(0))


def build_estimate(request, *, trusted_resolution=None):
    if request.get('price_basis') != 'BASE_NET':
        raise ValueError('Composition requires BASE_NET; VAT-inclusive components prohibited')
    mode = request.get('mode')
    if mode not in ('BASE_ONLY', 'PREVIEW'):
        raise ValueError('Only BASE_ONLY/PREVIEW construction is supported')
    source_lines = request.get('lines')
    if not isinstance(source_lines, list) or not source_lines:
        raise ValueError('At least one source-backed line required')
    missing = []; lines = []; seen = set()
    for item in source_lines:
        line_missing_start = len(missing)
        trace = item.get('trace_id')
        if not isinstance(trace, str) or not trace or trace in seen:
            raise ValueError('Unique line trace_id required')
        scope = item.get('scope', request.get('scope'))
        if not isinstance(scope, (str, dict)) or not scope:
            raise ValueError('Explicit composition scope required')
        seen.add(trace); norm = item.get('norm', {})
        for key in ('code', 'revision', 'unit', 'physical_unit', 'source'):
            if not isinstance(norm.get(key), str) or not norm[key]:
                raise ValueError('Norm requires explicit ' + key)
        evidence(norm.get('evidence'))
        quantity = dict(item.get('quantity', {})); quantity['trace_id'] = trace
        if quantity.get('physical_unit') != norm['physical_unit'] or decimal(quantity.get('norm_factor')) != decimal(norm.get('norm_factor')):
            raise ValueError('Quantity unit/factor must match vetted norm')
        converted = convert_quantity(quantity); amount = decimal(converted['norm_units'])
        if not resolved_authority(trusted_resolution, 'norm', trace, norm['evidence'],
                                  {'norm': authority_facts(norm), 'scope': scope, 'quantity': quantity,
                                   'work': {key: item.get(key) for key in ('description', 'technology', 'materials', 'technical_conditions')}}):
            missing.append(trace + ': norm applicability not resolved by server authority')
        components = {}
        for component in COMPONENTS:
            raw = norm.get('components', {}).get(component)
            if raw is None:
                missing.append(trace + ': missing norm component ' + component)
                components[component] = None
            else:
                components[component] = decimal(raw) * amount
        if components['machines'] is not None and components['machinist'] is not None and components['machinist'] > components['machines']:
            raise ValueError('Machinist component cannot exceed machines')
        resources = []; resource_ids = set()
        for resource in norm.get('resources', []):
            rid = resource.get('id')
            if not isinstance(rid, str) or not rid or rid in resource_ids:
                raise ValueError('Unique resource id required within norm')
            resource_ids.add(rid); evidence(resource.get('evidence'))
            if not isinstance(resource.get('unit'), str) or not resource['unit']:
                raise ValueError('Resource unit required')
            if type(resource.get('included')) is not bool:
                raise ValueError('Explicit included/notcount resource flag required')
            consumption = resource.get('quantity_per_norm')
            project_quantity = resource.get('project_quantity')
            if project_quantity is not None:
                if project_quantity.get('confirmed') is not True:
                    raise ValueError('Project resource quantity requires confirmation')
                evidence(project_quantity.get('evidence'))
                if project_quantity.get('unit') != resource['unit']:
                    raise ValueError('Project resource quantity unit must match resource')
                consumed = decimal(project_quantity.get('value'))
            elif consumption is None or (isinstance(consumption, str) and consumption.strip().casefold() in ('п', 'project', 'unknown')):
                consumed = None
                missing.append(f'{trace}: resource {rid} quantity is project-defined/unknown; confirmed project consumption required')
            else:
                consumed = decimal(consumption) * amount
            external = resource.get('external_price')
            if resource['included'] and external is not None:
                raise ValueError('Included resource already priced: double counting rejected')
            cost = None
            if not resource['included']:
                if consumed is None:
                    components['materials'] = None
                if external is None:
                    missing.append(f'{trace}: excluded resource {rid} price missing')
                    components['materials'] = None
                else:
                    if external.get('basis') != 'BASE_NET' or external.get('unit') != resource['unit']:
                        raise ValueError('Excluded resource price basis/unit must match BASE_NET resource')
                    evidence(external.get('evidence')); price = decimal(external.get('value'))
                    cost = None if consumed is None else consumed * price
                    if components['materials'] is not None and cost is not None:
                        components['materials'] += cost
            resources.append({'id': rid, 'unit': resource['unit'], 'included': resource['included'],
                              'quantity': str(consumed) if consumed is not None else None,
                              'quantity_origin': 'CONFIRMED_PROJECT_INPUT' if project_quantity is not None else 'NORM_PER_UNIT' if consumed is not None else 'UNKNOWN_PROJECT_DEFINED',
                              'project_quantity_evidence': project_quantity.get('evidence') if project_quantity is not None else None,
                              'external_cost_base': rounded(cost), 'evidence': resource['evidence']})
        coefficients = item.get('coefficients', [])
        if not coefficients:
            no_extra = item.get('no_extra_coefficients')
            if not isinstance(no_extra, dict) or not no_extra.get('evidence') or not resolved_authority(
                    trusted_resolution, 'no_extra_coefficients', trace, no_extra.get('evidence'),
                    {'determination': authority_facts(no_extra), 'norm_identity': {k: norm[k] for k in ('code', 'revision', 'source')}, 'scope': scope}):
                missing.append(trace + ': explicit no-extra-coefficients determination not resolved by server authority')
        adjusted = apply_factors(components, coefficients, trace, missing, resolver=trusted_resolution, scope=scope)
        if adjusted['machines'] is not None and adjusted['machinist'] is not None and adjusted['machinist'] > adjusted['machines']:
            raise ValueError('Scoped factors leave machinist greater than machines')
        priced = adjusted
        indexes = item.get('indexes', [])
        if mode == 'PREVIEW':
            index_components = [v.get('component') for v in indexes]
            if len(index_components) != len(set(index_components)):
                raise ValueError('Duplicate index component: index applied twice')
            if set(index_components) != set(COMPONENTS):
                missing.append(trace + ': indexes required for all four components')
                priced = {key: None for key in COMPONENTS}
            else:
                priced = apply_factors(adjusted, indexes, trace, missing, resolver=trusted_resolution, kind='index', scope=scope)
        elif indexes:
            raise ValueError('BASE_ONLY cannot apply current indexes')
        labour = priced['labour']; machinist = priced['machinist']
        fot = None if labour is None or machinist is None else labour + machinist
        extras = {}
        for field in ('overhead', 'profit'):
            rule = item.get(field)
            if rule is None:
                missing.append(trace + ': missing explicit ' + field + ' rate')
                extras[field] = None
            else:
                if rule.get('basis') != 'FOT' or rule.get('scope_trace_id') != trace:
                    raise ValueError('NR/SP require explicit FOT basis and matching scope')
                evidence(rule.get('evidence'))
                if not resolved_authority(trusted_resolution, field, trace, rule['evidence'], {'rule': authority_facts(rule), 'scope': scope}):
                    missing.append(trace + ': ' + field + ' applicability not resolved by server authority')
                extras[field] = None if fot is None else fot * decimal(rule.get('percent')) / 100
        current_blocked = mode == 'PREVIEW' and len(missing) > line_missing_start
        if current_blocked:
            priced = {k: None for k in COMPONENTS}; fot = None
            extras = {'overhead': None, 'profit': None}
        direct = total(priced)
        line_total_exact = None if direct is None or any(v is None for v in extras.values()) else direct + extras['overhead'] + extras['profit']
        # The payable display is reconciled to displayed direct/NR/SP amounts.
        line_total = None if line_total_exact is None else sum((Decimal(rounded(v)) for v in (direct, extras['overhead'], extras['profit'])), Decimal(0))
        lines.append({'trace_id': trace, 'scope': scope, 'description': item.get('description', norm['code']), 'norm': {k: norm[k] for k in ('code', 'revision', 'unit', 'source', 'evidence')},
                      'quantity': converted, 'resources': resources, 'components_base_exact': {k: str(v) if v is not None else None for k, v in adjusted.items()},
                      'components_exact': {k: str(v) if v is not None else None for k, v in priced.items()},
                      'direct': rounded(direct), 'fot': rounded(fot), 'overhead': rounded(extras['overhead']), 'profit': rounded(extras['profit']),
                      'current_costs_blocked': current_blocked,
                      'net': rounded(line_total), 'net_exact': str(line_total_exact) if line_total_exact is not None else None})
    net = None if any(l['net'] is None for l in lines) else sum((Decimal(l['net']) for l in lines), Decimal(0))
    diagnostic_net = None if any(l['net_exact'] is None for l in lines) else sum((Decimal(l['net_exact']) for l in lines), Decimal(0))
    vat_rule = request.get('vat'); vat = None; gross = None
    if vat_rule is None:
        missing.append('VAT regime/rate not supplied; gross total unavailable')
    else:
        if vat_rule.get('already_included') is not False:
            raise ValueError('VAT already included/unspecified: double VAT rejected')
        evidence(vat_rule.get('evidence'))
        vat_authorized = resolved_authority(trusted_resolution, 'vat', '__estimate__', vat_rule['evidence'],
                                           {'rule': authority_facts(vat_rule), 'scope': request.get('scope')})
        if not vat_authorized:
            missing.append('VAT applicability not resolved by server authority')
        vat_exact = None if net is None else net * decimal(vat_rule.get('percent')) / 100
        vat = None if vat_exact is None or (mode == 'PREVIEW' and not vat_authorized) else Decimal(rounded(vat_exact))
        gross = None if net is None or vat is None else net + vat
    return {'schema': 'estimate-calculation-v1', 'executor': 'DETERMINISTIC_TOOL', 'status': mode + '_DRAFT',
            'approval': 'NOT_FOR_APPROVAL', 'scenario_only': True, 'normative_applicability': 'UNCONFIRMED' if missing else 'SERVER_RESOLVED_FOR_DRAFT',
            'lines': lines, 'totals': {'net': rounded(net), 'vat': rounded(vat), 'gross': rounded(gross)},
            'diagnostics': {'unrounded_net_exact': str(diagnostic_net) if diagnostic_net is not None else None},
            'missing_data': missing, 'rounding': 'Half-up 0.01 direct/NR/SP; net sums displayed amounts; estimate net sums displayed lines; VAT rounds from displayed net; exact diagnostics are not payable totals',
            'native_grand_export': 'NOT_SUPPORTED; JSON/XLSX review draft only'}


def cost_effect(request):
    """Signed scenario, never a substituted or approved estimate total."""
    conditions = request.get('conditions')
    if not isinstance(conditions, list) or not conditions or any(not isinstance(x, str) or not x for x in conditions):
        raise ValueError('Explicit nonempty scenario conditions required')
    for key in ('trace_id', 'resource_id', 'unit', 'source'):
        if not isinstance(request.get(key), str) or not request[key]:
            raise ValueError('Cost effect requires ' + key)
    if not isinstance(request.get('scope'), str) or not request['scope']:
        raise ValueError('Explicit scope required; resource-only scenario is not a full estimate')
    if request.get('same_resource') is not True or request.get('old_unit') != request['unit'] or request.get('new_unit') != request['unit']:
        raise ValueError('Same resource and matching old/new units required')
    missing = []; delta = None; components = {}; effect = None
    for key in ('old_quantity', 'new_quantity'):
        if request.get(key) is None:
            missing.append(key + ' missing')
    if not missing:
        delta = decimal(request['new_quantity']) - decimal(request['old_quantity'])
    prices = request.get('per_unit_components')
    if not isinstance(prices, dict) or not prices:
        missing.append('exact source per-unit components/price missing')
    else:
        if request.get('price_basis') not in ('BASE_NET', 'CURRENT_NET'):
            raise ValueError('Scenario requires explicit net price basis')
        evidence(request.get('price_evidence'))
        for key, raw in prices.items():
            if key not in ('labour', 'machines', 'materials', 'overhead', 'profit'):
                raise ValueError('Unknown/additive component; machinist is included in machines')
            if raw is None:
                missing.append('per-unit ' + key + ' missing'); components[key] = None
            else:
                value = decimal(raw); components[key] = str(value * delta) if delta is not None else None
        if delta is not None and not missing:
            effect = sum((Decimal(v) for v in components.values()), Decimal(0))
    return {'schema': 'estimate-cost-effect-v1', 'executor': 'DETERMINISTIC_TOOL', 'status': 'CONDITIONAL_SCENARIO',
            'approval': 'NOT_FOR_APPROVAL', 'trace_id': request['trace_id'], 'resource_id': request['resource_id'], 'unit': request['unit'],
            'quantity_delta': str(delta) if delta is not None else None, 'component_effects_exact': components,
            'cost_effect_net': rounded(effect), 'scope': request['scope'],
            'conditions': conditions, 'source': request['source'], 'price_evidence': request.get('price_evidence'), 'missing_data': missing}


def export_xlsx(result, path):
    """Portable review worksheet, not a signed or native GRAND-Smeta document."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book = Workbook(); sheet = book.active; sheet.title = 'Проверочный расчёт'
    def numeric(value):
        return decimal(value, nonnegative=False) if value is not None else 'НЕИЗВЕСТНО'
    sheet.append(['ЧЕРНОВИК ДЛЯ ПРОФЕССИОНАЛЬНОЙ ПРОВЕРКИ — НЕ УТВЕРЖДЕНО'])
    sheet.append(['Статус', result['status']]); sheet.append(['Исполнитель', 'DETERMINISTIC_TOOL'])
    if 'lines' in result:
        sheet.append(['Trace ID', 'Работа', 'Норматив', 'Редакция', 'Физический объём', 'Ед.', 'Нормативный объём', 'Прямые', 'НР', 'СП', 'Без НДС', 'Источник'])
        for line in result['lines']:
            q = line['quantity']; n = line['norm']
            sheet.append([line['trace_id'], line['description'], n['code'], n['revision'], numeric(q['physical']), q['physical_unit'], numeric(q['norm_units']), numeric(line['direct']), numeric(line['overhead']), numeric(line['profit']), numeric(line['net']), n['source']])
            for column in (5, 7): sheet.cell(sheet.max_row, column).number_format = '0.########'
            for column in (8, 9, 10, 11): sheet.cell(sheet.max_row, column).number_format = '0.00'
        for key, value in result['totals'].items():
            sheet.append(['Итог сценария: ' + key, numeric(value)])
            sheet.cell(sheet.max_row, 2).number_format = '0.00'
    else:
        for key in ('trace_id', 'resource_id', 'unit', 'quantity_delta', 'cost_effect_net', 'scope', 'source'):
            is_numeric = key in ('quantity_delta', 'cost_effect_net')
            value = numeric(result.get(key)) if is_numeric else result.get(key) if result.get(key) is not None else 'НЕИЗВЕСТНО'
            sheet.append([key, value])
            if is_numeric: sheet.cell(sheet.max_row, 2).number_format = '0.00' if key == 'cost_effect_net' else '0.########'
        for condition in result['conditions']: sheet.append(['Условие сценария', condition])
    for issue in result.get('missing_data', []): sheet.append(['Открытый вопрос', issue])
    # Untrusted text must not become a spreadsheet formula/link executable.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str): cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    sheet['A1'].font = Font(bold=True, color='9C0006'); sheet.freeze_panes = 'A5'
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    book.save(Path(path)); book.close()


def main():
    parser = argparse.ArgumentParser(); parser.add_argument('input', type=Path); parser.add_argument('--operation', choices=('convert', 'compose', 'effect'), required=True)
    parser.add_argument('--output', type=Path, required=True); parser.add_argument('--xlsx', type=Path)
    args = parser.parse_args()
    from estimate_storage import authorized_input, managed_output, register_output
    args.input=authorized_input(args.input)
    args.output=managed_output(args.output,'estimate-calculation:'+str(args.output.absolute()))
    if args.xlsx:args.xlsx=managed_output(args.xlsx,'estimate-calculation-xlsx:'+str(args.xlsx.absolute()))
    if args.input.stat().st_size > 1024 * 1024:
        raise ValueError('Input request budget exceeded')
    request = json.loads(args.input.read_text())
    result = {'convert': convert_quantity, 'compose': build_estimate, 'effect': cost_effect}[args.operation](request)
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    register_output(args.output,'estimate-calculation:'+str(args.output),derived_from=str(args.input))
    if args.xlsx:
        if args.operation == 'convert': raise ValueError('XLSX export requires composition/scenario')
        export_xlsx(result, args.xlsx)
        register_output(args.xlsx,'estimate-calculation-xlsx:'+str(args.xlsx),derived_from=str(args.output))
    print(json.dumps({'output': str(args.output), 'operation': args.operation}))


if __name__ == '__main__': main()
