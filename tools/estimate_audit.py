"""Read-only estimate inspection. All numbers use Decimal, never LLM arithmetic.

No normative approval, source modification, GSFX signing or database promotion.
"""
import ast
import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from decimal import Decimal, InvalidOperation, DecimalException, ROUND_HALF_UP
from pathlib import Path, PurePosixPath

MAX_BYTES = 100 * 1024 * 1024


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


def number(value):
    if value is None or isinstance(value, bool):
        raise ValueError('Unknown or boolean is not a number')
    result = Decimal(str(value).replace('\u00a0', '').replace(' ', '').replace(',', '.'))
    if not result.is_finite():
        raise ValueError('Nonfinite number')
    return result


def arithmetic(expression, variables=None):
    """Bounded arithmetic only; no eval, names/functions/cells unless supplied."""
    text = str(expression).strip().lstrip('=').replace(',', '.').replace('×', '*').replace('−', '-')
    rounded=re.fullmatch(r'ОКР\((.*);\s*(\d{1,2})\)',text,re.I)
    if rounded:
        digits=int(rounded[2])
        if digits>12:raise ValueError('Unsupported precision')
        return arithmetic(rounded[1],variables).quantize(Decimal(1).scaleb(-digits),rounding=ROUND_HALF_UP)
    if len(text) > 1000:
        raise ValueError('Expression too long')
    node = ast.parse(text, mode='eval')
    if sum(1 for _ in ast.walk(node)) > 100:
        raise ValueError('Expression too complex')
    variables = variables or {}
    def calc(n):
        if isinstance(n, ast.Constant) and type(n.value) in (int, float):
            return number(ast.get_source_segment(text, n))
        if isinstance(n, ast.Name) and n.id in variables:
            return number(variables[n.id])
        if isinstance(n, ast.UnaryOp) and isinstance(n.op, (ast.UAdd, ast.USub)):
            return calc(n.operand) * (-1 if isinstance(n.op, ast.USub) else 1)
        if isinstance(n, ast.BinOp) and isinstance(n.op, (ast.Add, ast.Sub, ast.Mult, ast.Div)):
            a, b = calc(n.left), calc(n.right)
            if isinstance(n.op, ast.Add): return a + b
            if isinstance(n.op, ast.Sub): return a - b
            if isinstance(n.op, ast.Mult): return a * b
            if b == 0: raise ValueError('Division by zero')
            return a / b
        raise ValueError('Unsupported expression or unresolved variable')
    return calc(node.body)


def zip_members(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        total = 0
        if len(z.infolist()) > 1000: raise ValueError('Too many archive members')
        for info in z.infolist():
            name = PurePosixPath(info.filename.replace('\\', '/'))
            if name.is_absolute() or '..' in name.parts or ':' in str(name):
                raise ValueError('Unsafe archive member')
            if info.external_attr >> 16 & 0o170000 == 0o120000:
                raise ValueError('Archive symlink rejected')
            total += info.file_size
            if total > MAX_BYTES: raise ValueError('Expanded archive too large')
            if not info.is_dir(): yield str(name), z.read(info)


def xml_read(raw):
    from lxml import etree
    if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
        raise ValueError('XML entities prohibited')
    root = etree.fromstring(raw, etree.XMLParser(resolve_entities=False, no_network=True))
    tree = root.getroottree()
    reference_values={}
    for p in root.findall('.//Position'):
        ident=p.get('Identifier');q=p.find('Quantity')
        if ident and q is not None and q.get('Result') is not None:
            reference_values[ident.casefold()]=q.get('Result')
            for resource in p.findall('./Resources/*'):
                if resource.get('Identifier') and resource.get('Quantity'):
                    try:
                        quantity=number(resource.get('Quantity'))
                        quantity_flag={'Mat':'MatQty','Mch':'EmQty','Tzr':'OzpTz','Tzm':'ZpmTz'}.get(resource.tag)
                        for coefficient in p.findall('./Koefficients/K'):
                            if quantity_flag and quantity_flag in coefficient.get('Options','').split():
                                # Literal native flag affects this resource quantity; legality is separate.
                                if not coefficient.get('Value_PZ'):raise ValueError('Resource quantity coefficient value unresolved')
                                quantity*=number(coefficient.get('Value_PZ'))
                        reference_values[(ident+'.'+resource.get('Identifier')).casefold()]=str(number(q.get('Result'))*quantity)
                    except (ValueError,InvalidOperation):pass
    nodes = [{'path': tree.getpath(n), 'tag': n.tag, 'attributes': dict(n.attrib),
              'text': n.text.strip() if n.text and n.text.strip() else None}
             for n in root.iter() if isinstance(n.tag, str)]
    positions = []
    for p in root.findall('.//Position'):
        positions.append({'path': tree.getpath(p), 'attributes': dict(p.attrib),
                          'active': 'Inactive' not in p.get('Options', '').split(),
                          'children': [{'tag': n.tag, 'attributes': dict(n.attrib)} for n in p]})
    formulas = []
    for n in root.iter():
        if not isinstance(n.tag,str): continue
        for key in ('Fx','Formula','Quantity'):
            if key not in n.attrib: continue
            expr=n.get(key); record={'path':tree.getpath(n), 'field':key,'expression':expr,
                                    'result':n.get('Result'),'status':'UNRESOLVED'}
            try:
                resolved=expr
                for ref in sorted(reference_values,key=len,reverse=True):
                    resolved=re.sub(re.escape(ref)+r'(?![\w.])','('+str(reference_values[ref])+')',resolved,flags=re.I)
                val=arithmetic(resolved)
                if key=='Fx' and n.tag=='Quantity' and n.get('KUnit'):
                    if number(n.get('KUnit'))<=0:raise ValueError('KUnit must be positive')
                    val=val/number(n.get('KUnit'))
                if key=='Fx' and n.tag=='Quantity' and n.get('KMult'):
                    val=val*number(n.get('KMult'))
                if key=='Fx' and n.tag=='Quantity' and n.get('Precision'):
                    val=val.quantize(Decimal(1).scaleb(-int(n.get('Precision'))),rounding=ROUND_HALF_UP)
                record.update(calculated=str(val),status='EVALUATED',k_unit=n.get('KUnit'),k_mult=n.get('KMult'))
                if n.get('Result') is not None:
                    delta=val-number(n.get('Result'));record.update(delta=str(delta),status='MATCH' if abs(delta)<=Decimal('.000001') else 'MISMATCH')
            except (ValueError,SyntaxError,DecimalException,OverflowError,RecursionError): pass
            formulas.append(record)
    ids=[p['attributes'].get('SysID') for p in positions]
    return {'root':root.tag, 'root_attributes':dict(root.attrib),'nodes':nodes,
            'positions':positions,'formulas':formulas,
            'counts':{'positions':len(positions),'active':sum(p['active'] for p in positions),
                      'duplicate_position_sysids':[x for x,c in Counter(ids).items() if x and c>1]},
            'hierarchy':dict(Counter(n['path'].rsplit('/',1)[0] for n in nodes)),
            'tag_counts':dict(Counter(n['tag'] for n in nodes))}


def spreadsheet_read(raw):
    import openpyxl
    # Validate expansion bounds before openpyxl parses any member.
    for _ in zip_members(raw):pass
    wf=openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    wv=openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    result=[]
    for ws in wf:
        if ws.max_row*ws.max_column>250000 or ws.max_row>10000 or ws.max_column>1000:
            raise ValueError('Worksheet dimension budget exceeded')
        values=wv[ws.title]; cells=[];formulas=[]
        for row in ws:
            for c in row:
                if c.value is None: continue
                val=c.value if isinstance(c.value,(str,int,float,bool)) else str(c.value)
                cells.append({'cell':c.coordinate,'value':val,'data_type':c.data_type,
                              'cached':values[c.coordinate].value,'number_format':c.number_format})
                if c.data_type=='f':
                    record={'cell':c.coordinate,'formula':val,'cached':values[c.coordinate].value,'status':'UNRESOLVED'}
                    variables={m.group():values[m.group()].value for m in re.finditer(r'\b[A-Z]+[1-9][0-9]*\b',val)}
                    try:
                        calculated=arithmetic(val,variables);record['calculated']=str(calculated)
                        cache=record['cached']
                        if cache is None:record['status']='NO_CACHE'
                        else:
                            delta=calculated-number(cache);record.update(delta=str(delta),status='MATCH' if abs(delta)<=Decimal('.000001') else 'MISMATCH')
                    except (ValueError,SyntaxError,DecimalException,OverflowError,RecursionError):pass
                    formulas.append(record)
        result.append({'sheet':ws.title,'rows':ws.max_row,'columns':ws.max_column,
                       'print_area':str(ws.print_area),'merged':[str(x) for x in ws.merged_cells.ranges],
                       'hidden_rows':[i for i,d in ws.row_dimensions.items() if d.hidden],
                       'hidden_columns':[i for i,d in ws.column_dimensions.items() if d.hidden],
                       'cells':cells,'formulas':formulas})
    wf.close();wv.close();return result


def inspect_file(name,raw,depth=0):
    if depth>3 or len(raw)>MAX_BYTES: raise ValueError('Input bounds exceeded')
    result={'name':name,'sha256':sha(raw),'size':len(raw),'format':Path(name).suffix.lower()}
    suffix=result['format']
    if suffix in ('.gsfx','.zip'):
        members=[]
        for member,data in zip_members(raw):
            item=inspect_file(member,data,depth+1)
            members.append(item)
        result['members']=members
    elif suffix in ('.xml','.gge'):result['xml']=xml_read(raw)
    elif suffix=='.xlsx':result['sheets']=spreadsheet_read(raw)
    elif suffix=='.pdf':
        import pymupdf
        with pymupdf.open(stream=raw,filetype='pdf') as pdf:
            if len(pdf)>100:raise ValueError('PDF page budget exceeded')
            result['pages']=[]
            for i,p in enumerate(pdf):
                text=p.get_text();item={'page':i+1,'text':text,'images':len(p.get_images()),'method':'TEXT_LAYER'}
                if not text.strip() or (item['images'] and len(text.strip())<400):
                    import subprocess
                    try:
                        image=p.get_pixmap(matrix=pymupdf.Matrix(2,2))
                        if image.width*image.height>25000000:raise ValueError('PDF image budget exceeded')
                        r=subprocess.run(['tesseract','stdin','stdout','-l','rus+eng'],input=image.tobytes('png'),capture_output=True,check=True,timeout=45)
                        item.update(text=r.stdout.decode(),text_layer=text,method='OCR_UNVERIFIED')
                    except (OSError,ValueError,subprocess.SubprocessError):item['method']='OCR_UNAVAILABLE'
                result['pages'].append(item)
            result['extraction']='Text/OCR extraction; numerical OCR requires independent image verification'
    elif suffix in ('.txt',):
        try:result['text']=raw.decode('utf-8-sig')
        except UnicodeDecodeError:result['text']=raw.decode('cp1251')
    return result


def inspect_directory(path):
    path=Path(path)
    files=[]
    for p in sorted(path.rglob('*')):
        if p.is_symlink():raise ValueError('Source symlinks prohibited')
        if p.is_file():
            if p.stat().st_size>MAX_BYTES:raise ValueError('Input too large')
            files.append(inspect_file(str(p.relative_to(path)),p.read_bytes()))
    return {'schema':'estimate-inspection-v1','executor':'DETERMINISTIC_TOOL','files':files,
            'source_modified':False,'normative_applicability':'NOT_VERIFIED','native_import':'NOT_CHECKED'}


def inspection_summary(data):
    out=[]
    for f in data['files']:
        item={k:f[k] for k in ('name','sha256','format','size')}
        if 'sheets' in f:
            item['sheets']=[{'sheet':s['sheet'],'rows':s['rows'],'columns':s['columns'],
                            'formula_states':dict(Counter(x['status'] for x in s['formulas']))} for s in f['sheets']]
        if 'members' in f:
            item['members']=[m['name'] for m in f['members']]
            for m in f['members']:
                if m['name']=='Data.xml':item.update(xml_counts=m['xml']['counts'],document_type=m['xml']['root_attributes'].get('DocumentType'))
        if 'xml' in f:item['root']=f['xml']['root']
        out.append(item)
    return out


def controls(data):
    """Source contradictions and numerical comparisons, not normative verdicts."""
    findings=[];checks=[];unverified_candidates=[]
    def finding(kind,classification,evidence,description):
        findings.append({'id':'F'+str(len(findings)+1).zfill(3),'kind':kind,'classification':classification,
                         'evidence':evidence,'description':description,'normative_basis':None,
                         'cost_effect':{'value':None,'missing_data':['Confirmed corrected quantity and scope','Confirmed relationship between changed work and priced resource','Source prices/components, permitted coefficients and applicable indexes/NR/SP/VAT for claimed scope']},'confidence':'document_literal'})
    def children(f):
        yield f
        for m in f.get('members',[]):yield from children(m)
    # Physical linear quantities from source worksheets; pairs remain candidates.
    linear=[]
    for f in data['files']:
        for s in f.get('sheets',[]):
            rowmap={}
            for c in s['cells']:rowmap.setdefault(re.search(r'\d+',c['cell'])[0],[]).append(c)
            for row,cs in rowmap.items():
                unit=next((c for c in cs if isinstance(c['value'],str) and c['value'].strip().lower() in ('м.п.','пог.м','пог. м','м.п','м')),None)
                if unit is None:continue
                desc=next((c for c in cs if isinstance(c['value'],str) and len(c['value'])>30),None)
                qty=next((c for c in cs if c['cell'].rstrip('0123456789') in ('E','H') and type(c['value']) in (int,float)),None)
                if desc and qty:linear.append({'file':f['name'],'sheet':s['sheet'],'row':row,'description':desc['value'],'quantity':str(qty['value']),'unit':unit['value'],'cell':qty['cell']})
    for f in data['files']:
        for doc in children(f):
            xml=doc.get('xml')
            if xml:
                for formula in xml['formulas']:
                    if formula['status']=='MISMATCH':finding('FORMULA_CACHE','POTENTIAL_ERROR',{'file':f['name'],'member':doc['name'],**formula},'Evaluated formula differs from cached result')
                checks.append({'type':'XML_FORMULAS','file':f['name'],'member':doc['name'],
                               'states':dict(Counter(v['status'] for v in xml['formulas']))})
            for sheet in doc.get('sheets',[]):
                cells={c['cell']:c for c in sheet['cells']};rows={}
                for c in sheet['cells']:
                    rows.setdefault(int(re.search(r'\d+',c['cell'])[0]),[]).append(c)
                for row,cs in rows.items():
                    for c in cs:
                        if not isinstance(c['value'],str):continue
                        match=re.search(r'толщин\w*\s*(\d+(?:[.,]\d+)?)\s*(см|мм)',c['value'],re.I)
                        if not match:continue
                        thickness=number(match[1])/(100 if match[2].lower()=='см' else 1000)
                        for fc in cs:
                            if fc is c or not isinstance(fc['value'],str):continue
                            expression=fc['value']
                            # Literal geometry formula with one length multiplier only.
                            m=re.fullmatch(r'=?\s*(\d+(?:[.,]\d+)?)\s*\*\s*(0[.,]\d+)\s*',expression)
                            if m and number(m[2])!=thickness:
                                actual=arithmetic(expression);expected=number(m[1])*thickness
                                finding('THICKNESS_CONTRADICTION','INSUFFICIENT_EVIDENCE',
                                        {'file':f['name'],'sheet':sheet['sheet'],'description_cell':c['cell'],
                                         'formula_cell':fc['cell'],'description':c['value'],'formula':expression,
                                         'calculated':str(actual),'description_scenario':str(expected),'delta':str(expected-actual)},
                                        'Description thickness conflicts with formula; drawing decides correct thickness')
                checks.append({'type':'XLSX_FORMULAS','file':f['name'],'sheet':sheet['sheet'],
                               'states':dict(Counter(v['status'] for v in sheet['formulas']))})
                for v in sheet['formulas']:
                    if v['status']=='MISMATCH':finding('XLSX_CACHE','POTENTIAL_ERROR',{'file':f['name'],'sheet':sheet['sheet'],**v},'XLSX formula differs from cache')
            if xml and doc['name']=='Data.xml':
                repeated={}
                for p in xml['positions']:
                    a=p['attributes'];link=a.get('NotCountedResSysID')
                    if link and p['active']:repeated.setdefault(link,[]).append({'path':p['path'],**a})
                    dim=re.search(r'(\d+(?:[.,]\d+)?)\s*[хx×]\s*\d+\s*[хx×]\s*\d+\s*мм',a.get('Caption',''),re.I)
                    q=next((n['attributes'].get('Result') for n in xml['nodes'] if n['path']==p['path']+'/Quantity'),None)
                    if p['active'] and a.get('Units')=='шт' and dim and q and number(q)>0:
                        tokens=set(re.findall(r'[а-я]{4,}',a.get('Caption','').lower()))
                        for source in linear:
                            if len(tokens & set(re.findall(r'[а-я]{4,}',source['description'].lower())))<2:continue
                            physical=number(source['quantity'])
                            implied=physical/number(q)
                            if True:
                                unverified_candidates.append({'kind':'LINEAR_PIECE_CONVERSION','status':'UNVERIFIED_MATCH_CANDIDATE','evidence':
                                        {'file':f['name'],'path':p['path'],'position':a.get('Number'),'caption':a['Caption'],
                                         'pieces':q,'source':source,'dimension_semantics':'UNCONFIRMED',
                                         'candidate_implied_piece_length_m':str(implied),'relationship_confirmed':False,
                                         'missing_data':['Confirmed link between source work and product position','Product passport identifying length axis and installed effective length']},
                                        'description':'Text similarity identifies a search candidate only; no position relationship or quantity error established'})
                for link,positions in repeated.items():
                    if len(positions)>1:unverified_candidates.append({'kind':'REUSED_RESOURCE_LINK','status':'ALLOCATION_NOT_VERIFIED','evidence':{'file':f['name'],'link':link,'positions':positions},'description':'Repeated internal resource link is compatible with split allocation; an error requires a scoped resource-balance check'})
                counterparties=[n for n in xml['nodes'] if n['tag']=='Counterparty']
                offers=[n for n in xml['nodes'] if n['tag']=='Item' and 'Supplier' in n['attributes']]
                for offer in offers:
                    a=offer['attributes'];index=int(a.get('Supplier','0'))-1
                    if index<0 or index>=len(counterparties):
                        finding('SUPPLIER_LINK','CONFIRMED_ERROR',{'file':f['name'],'path':offer['path'],'supplier':a.get('Supplier')},'Supplier index cannot resolve');continue
                    party=counterparties[index]['attributes']
                    for pdf in f.get('members',[]):
                        if pdf['format']!='.pdf':continue
                        # Binding from the actual hyperlink, not similarity of supplier name.
                        link=[n for n in xml['nodes'] if n['path'].startswith(offer['path']+'/Hyperlinks/') and n['attributes'].get('Caption')==Path(pdf['name']).name]
                        if not link:continue
                        for page in pdf.get('pages',[]):
                            found=None
                            company=re.sub(r'[^а-яa-z0-9]','',party['Caption'].lower().replace('ооо',''))
                            for candidate in re.finditer(r'ИНН(?:/КПП)?\s*[:/]?\s*(\d{10,12})',page['text'],re.I):
                                prefix=re.sub(r'[^а-яa-z0-9]','',page['text'][:candidate.start()].lower())
                                if company and company in prefix:
                                    tail=page['text'][max(0,candidate.start()-70):candidate.start()]
                                    if not re.search(r'покупател|получател|заказчик',tail,re.I):found=candidate;break
                            if found and found[1]!=party.get('INN'):
                                finding('SUPPLIER_INN','CONFIRMED_ERROR' if page['method']=='TEXT_LAYER' else 'POTENTIAL_ERROR',
                                        {'file':f['name'],'path':offer['path'],'party':party['Caption'],
                                         'xml_inn':party.get('INN'),'pdf_inn':found[1],'pdf':pdf['name'],'page':page['page'],'method':page['method']},
                                        'Bound offer PDF and XML supplier INN disagree; correct identity requires confirmation')
                    try:
                        net=number(a['OptPrice']);gross=number(a['OptPriceWithVAT']);vat=number(a['VAT'])
                        delta=(net*(1+vat/100)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)-gross
                        checks.append({'type':'OFFER_VAT','file':f['name'],'path':offer['path'],'delta':str(delta),'status':'MATCH' if abs(delta)<=Decimal('.01') else 'MISMATCH'})
                    except (KeyError,ValueError,DecimalException):pass
                for n in xml['nodes']:
                    if '???' in json.dumps(n['attributes'],ensure_ascii=False):
                        finding('UNFINISHED_REQUISITE','INSUFFICIENT_EVIDENCE',{'file':f['name'],'path':n['path'],'attributes':n['attributes']},'Unfilled template requisite retained')
    return {'executor':'DETERMINISTIC_TOOL','findings':findings,'checks':checks,'unverified_candidates':unverified_candidates,
            'native_import':'NOT_CHECKED','normative_applicability':'NOT_VERIFIED'}


if __name__=='__main__':
    import argparse
    from estimate_storage import authorized_input,managed_output,register_output
    p=argparse.ArgumentParser();p.add_argument('directory');p.add_argument('--output',required=True);a=p.parse_args()
    # The directory reader rejects links itself; enforce managed ingestion scope.
    import importlib.util
    spec=importlib.util.spec_from_file_location('estimate_storage_policy','/opt/ai-station/storage.py')
    service=importlib.util.module_from_spec(spec);spec.loader.exec_module(service)
    service.ingestion_allowed(Path(a.directory));service.resolve_storage('estimates','staging')
    output=managed_output(a.output,'estimate-inspection:'+str(Path(a.output).absolute()))
    data=inspect_directory(a.directory)
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str))
    register_output(output,'estimate-inspection:'+str(output))
    print(json.dumps({'files':len(data['files']),'output':str(output)}))
